import os
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")

import django  # noqa: E402

django.setup()

from django.apps import apps  # noqa: E402
from django.db.migrations.recorder import MigrationRecorder  # noqa: E402

HEADERS = [
    "Имя поля",
    "Описание",
    "Тип данных",
    "Ключевое поле",
    "Допускается пустой ввод",
]

TABLE_DESCRIPTIONS = {
    "api_department": "Подразделения предприятия.",
    "api_userprofile": "Профили сотрудников с ролью и привязкой к отделу.",
    "api_equipment": "Оборудование, его состояние и срок годности.",
    "api_task": "Производственные задачи и назначения.",
    "api_comment": "Комментарии к задачам.",
    "api_incident": "Инциденты по безопасности и поломкам.",
    "api_message": "Внутренние сообщения между сотрудниками.",
    "auth_user": "Учетные записи пользователей Django.",
    "auth_group": "Группы прав доступа.",
    "auth_permission": "Права доступа к моделям и действиям.",
    "auth_user_groups": "Связь пользователей и групп (M2M).",
    "auth_user_user_permissions": "Пользовательские права напрямую (M2M).",
    "auth_group_permissions": "Связь групп и прав (M2M).",
    "django_content_type": "Реестр моделей Django для generic-механизмов.",
    "django_admin_log": "Журнал действий в админ-панели.",
    "django_session": "Сессионные данные пользователей.",
    "django_migrations": "История примененных миграций.",
}

FIELD_DESCRIPTIONS = {
    "id": "Уникальный идентификатор записи.",
    "name": "Название сущности.",
    "created_at": "Дата и время создания записи.",
    "updated_at": "Дата и время последнего изменения записи.",
    "title": "Краткий заголовок.",
    "description": "Развернутое текстовое описание.",
    "status": "Текущий статус записи.",
    "type": "Тип события/сущности.",
    "text": "Текстовое содержимое.",
    "read": "Признак, что сообщение прочитано.",
    "role": "Роль сотрудника в системе.",
    "expiration_date": "Дата истечения срока годности/эксплуатации.",
    "photo": "Путь к файлу фотографии в хранилище media.",
    "due_date": "Плановый срок выполнения.",
    "duration_hours": "Оценка длительности задачи в часах.",
    "urgency": "Приоритет инцидента.",
    "username": "Уникальное имя для входа.",
    "first_name": "Имя пользователя.",
    "last_name": "Фамилия пользователя.",
    "email": "Email пользователя.",
    "password": "Хеш пароля.",
    "is_staff": "Доступ к административной панели.",
    "is_active": "Признак активной учетной записи.",
    "is_superuser": "Признак суперпользователя.",
    "date_joined": "Дата регистрации учетной записи.",
    "last_login": "Дата и время последнего входа.",
    "app_label": "Имя Django-приложения.",
    "model": "Имя модели в нижнем регистре.",
    "session_key": "Уникальный ключ сессии.",
    "session_data": "Сериализованные данные сессии.",
    "expire_date": "Дата и время истечения сессии.",
    "codename": "Системный код права доступа.",
    "action_time": "Момент фиксации действия.",
    "object_id": "ID объекта, над которым выполнено действие.",
    "object_repr": "Строковое представление объекта.",
    "action_flag": "Код типа действия (add/change/delete).",
    "change_message": "Описание изменений.",
    "sender": "Пользователь-отправитель.",
    "receiver": "Пользователь-получатель.",
    "author": "Пользователь-автор записи.",
    "assignee": "Назначенный исполнитель задачи.",
    "creator": "Пользователь, создавший запись.",
    "reporter": "Пользователь, сообщивший об инциденте.",
    "department": "Подразделение, к которому относится запись.",
    "equipment": "Оборудование, связанное с инцидентом.",
    "task": "Задача, к которой относится комментарий.",
    "user": "Связанный пользователь.",
    "group": "Связанная группа прав.",
    "permission": "Связанное право доступа.",
    "content_type": "Тип контента (модель) для generic-ссылок.",
    "user_id": "Идентификатор пользователя.",
    "group_id": "Идентификатор группы.",
    "permission_id": "Идентификатор права доступа.",
    "department_id": "Идентификатор подразделения.",
    "assignee_id": "Идентификатор исполнителя.",
    "creator_id": "Идентификатор автора записи.",
    "author_id": "Идентификатор автора.",
    "reporter_id": "Идентификатор инициатора инцидента.",
    "equipment_id": "Идентификатор оборудования.",
    "task_id": "Идентификатор задачи.",
    "sender_id": "Идентификатор отправителя.",
    "receiver_id": "Идентификатор получателя.",
    "content_type_id": "Идентификатор типа контента.",
    "user_name": "Имя миграции/пользователя (служебное поле).",
    "app": "Имя приложения, к которому относится миграция.",
    "migration": "Имя миграции.",
    "applied": "Дата и время применения миграции.",
}

CHOICE_LABELS = {
    "api_userprofile.role": "Допустимые значения: ADMIN, DEPARTMENT_HEAD, WORKER.",
    "api_equipment.status": "Допустимые значения: OPERATIONAL, BROKEN, EXPIRED.",
    "api_task.status": "Допустимые значения: PENDING, COMPLETED.",
    "api_incident.type": "Допустимые значения: ACCIDENT, BROKEN_EQUIPMENT.",
    "api_incident.status": "Допустимые значения: OPEN, RESOLVED.",
    "api_incident.urgency": "Допустимые значения: HIGH, CRITICAL.",
}

COMPOSITE_UNIQUES = {
    "auth_permission": ["content_type_id + codename"],
    "django_content_type": ["app_label + model"],
    "auth_user_groups": ["user_id + group_id"],
    "auth_user_user_permissions": ["user_id + permission_id"],
    "auth_group_permissions": ["group_id + permission_id"],
}


def humanize_type(field):
    internal = field.get_internal_type()
    if internal == "BigAutoField":
        return "bigint (auto increment)"
    if internal == "AutoField":
        return "integer (auto increment)"
    if internal == "BigIntegerField":
        return "bigint"
    if internal in {"IntegerField", "PositiveIntegerField", "PositiveSmallIntegerField", "SmallIntegerField"}:
        return "integer"
    if internal == "BooleanField":
        return "boolean"
    if internal == "DateTimeField":
        return "timestamp with time zone"
    if internal == "DateField":
        return "date"
    if internal == "TextField":
        return "text"
    if internal in {"CharField", "SlugField", "FileField"}:
        max_length = getattr(field, "max_length", None)
        return f"varchar({max_length})" if max_length else "varchar"
    if internal == "JSONField":
        return "jsonb"
    if internal == "EmailField":
        max_length = getattr(field, "max_length", 254)
        return f"varchar({max_length})"
    return internal


def key_info(model, field):
    chunks = []
    if field.primary_key:
        chunks.append("PK")

    if getattr(field, "remote_field", None) and field.many_to_one:
        ref = field.remote_field.model._meta.db_table
        ref_pk = field.remote_field.model._meta.pk.column
        chunks.append(f"FK -> {ref}.{ref_pk}")

    if field.unique and not field.primary_key:
        chunks.append("UNIQUE")

    unique_sets = COMPOSITE_UNIQUES.get(model._meta.db_table, [])
    for unique_def in unique_sets:
        if field.column in unique_def:
            chunks.append(f"UNIQUE({unique_def})")

    return "; ".join(chunks) if chunks else "-"


def nullable_info(field):
    return "Да" if (getattr(field, "null", False) or getattr(field, "blank", False)) else "Нет"


def field_description(table_name, field):
    base = FIELD_DESCRIPTIONS.get(field.column) or FIELD_DESCRIPTIONS.get(field.name) or "Поле данных таблицы."
    choice_note = CHOICE_LABELS.get(f"{table_name}.{field.name}")
    if choice_note:
        return f"{base} {choice_note}"
    return base


def collect_models():
    all_models = list(apps.get_models(include_auto_created=True))
    all_models.append(MigrationRecorder.Migration)

    unique_models = {}
    for model in all_models:
        table = model._meta.db_table
        if table not in unique_models:
            unique_models[table] = model

    return dict(sorted(unique_models.items(), key=lambda x: x[0]))


def write_workbook(output_path):
    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="1F1F1F")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    models_map = collect_models()

    for table_name, model in models_map.items():
        ws = wb.create_sheet(title=table_name[:31])

        ws["A1"] = f"Таблица: {table_name}"
        ws["A2"] = f"Описание: {TABLE_DESCRIPTIONS.get(table_name, 'Служебная таблица Django или прикладная таблица проекта.')}"
        ws.merge_cells("A1:E1")
        ws.merge_cells("A2:E2")
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"].font = Font(italic=True, color="555555")

        for col_idx, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=4, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        fields = [f for f in model._meta.concrete_fields]

        row = 5
        for field in fields:
            ws.cell(row=row, column=1, value=field.column)
            ws.cell(row=row, column=2, value=field_description(table_name, field))
            ws.cell(row=row, column=3, value=humanize_type(field))
            ws.cell(row=row, column=4, value=key_info(model, field))
            ws.cell(row=row, column=5, value=nullable_info(field))

            for c in range(1, 6):
                ws.cell(row=row, column=c).alignment = Alignment(vertical="top", wrap_text=True)
                ws.cell(row=row, column=c).border = border

            row += 1

        ws.freeze_panes = "A5"
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 70
        ws.column_dimensions["C"].width = 28
        ws.column_dimensions["D"].width = 40
        ws.column_dimensions["E"].width = 24

    wb.save(output_path)


if __name__ == "__main__":
    project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    output_file = os.path.join(project_root, "factoryproject_db_dictionary.xlsx")
    write_workbook(output_file)
    print(output_file)
